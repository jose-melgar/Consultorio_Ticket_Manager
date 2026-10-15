"""SQLite is authoritative; legacy files are imported once, never overwritten."""
from contextlib import closing, contextmanager
from datetime import datetime
import hashlib
import json
from pathlib import Path
import re
import shutil
import sqlite3
import threading
import uuid

from models import ValidationError, money, now


class StorageError(RuntimeError):
    pass


class ConflictError(ValidationError):
    pass


def encode(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, allow_nan=False)


def request_hash(value):
    try:
        return hashlib.sha256(encode(value).encode('utf-8')).hexdigest()
    except (ValueError, TypeError):
        raise ValidationError('La solicitud contiene datos JSON inválidos.') from None


class Store:
    def __init__(self, data_dir, tickets_dir):
        self.data_dir, self.tickets_dir = Path(data_dir), Path(tickets_dir)
        self.path = self.data_dir / 'consultorio.sqlite3'
        self._lock = threading.Lock()
        self._ready = False

    @contextmanager
    def connection(self):
        db = sqlite3.connect(self.path, timeout=30)
        db.row_factory = sqlite3.Row
        db.execute('PRAGMA foreign_keys=ON')
        db.execute('PRAGMA synchronous=FULL')
        try:
            yield db
            db.commit()
        except Exception:
            db.rollback()
            raise
        finally:
            db.close()

    def initialize(self, recover_jobs=True):
        with self._lock:
            if self._ready:
                return
            self.data_dir.mkdir(parents=True, exist_ok=True)
            with self.connection() as db:
                db.executescript('''
                    CREATE TABLE IF NOT EXISTS metadata(key TEXT PRIMARY KEY, value TEXT NOT NULL);
                    CREATE TABLE IF NOT EXISTS counters(name TEXT PRIMARY KEY, value INTEGER NOT NULL CHECK(value>=0));
                    CREATE TABLE IF NOT EXISTS tickets(
                        id TEXT PRIMARY KEY, specific_id TEXT UNIQUE NOT NULL,
                        operation_key TEXT UNIQUE NOT NULL, request_hash TEXT NOT NULL,
                        payload TEXT NOT NULL, created_at TEXT NOT NULL);
                    CREATE TABLE IF NOT EXISTS events(
                        id INTEGER PRIMARY KEY, ticket_id TEXT REFERENCES tickets(id),
                        kind TEXT NOT NULL, at TEXT NOT NULL, details TEXT NOT NULL);
                    CREATE TABLE IF NOT EXISTS jobs(
                        id INTEGER PRIMARY KEY, ticket_id TEXT NOT NULL REFERENCES tickets(id),
                        kind TEXT NOT NULL CHECK(kind IN ('excel','pdf','print')),
                        state TEXT NOT NULL DEFAULT 'pending', error TEXT);
                    CREATE INDEX IF NOT EXISTS jobs_state ON jobs(state,id);
                ''')
                db.execute('BEGIN IMMEDIATE')
                version = db.execute("SELECT value FROM metadata WHERE key='schema_version'").fetchone()
                if version and version['value'] != '1':
                    raise StorageError('La base requiere otra versión del programa.')
                if not version:
                    self._import_legacy(db)
                    db.execute("INSERT INTO metadata VALUES ('schema_version','1')")
                # A print interrupted after spool submission must not be repeated automatically.
                if recover_jobs:
                    db.execute("UPDATE jobs SET state='uncertain',error='Impresión interrumpida; compruebe el papel antes de reimprimir.' WHERE state='running' AND kind='print'")
                    db.execute("UPDATE jobs SET state='pending' WHERE state='running' AND kind!='print'")
            self._ready = True

    def _import_legacy(self, db):
        def read_json(name, default):
            path = self.data_dir / name
            if not path.exists():
                return default
            try:
                return json.loads(path.read_text(encoding='utf-8-sig'))
            except (ValueError, OSError) as exc:
                raise StorageError(f'No se puede importar {name}. Restaure una copia válida; no se modificó el original.') from exc
        history = read_json('ventas.json', [])
        counters = read_json('id_counters.json', {})
        if not isinstance(history, list) or not isinstance(counters, dict):
            raise StorageError('Formato de datos heredados inválido.')
        maximum = {'global': 0, 'general': 0, 'dental': 0}
        patterns = {'TK': 'global', 'LM': 'general', 'DEN': 'dental'}

        def reserve(value):
            for prefix, digits in re.findall(r'\b(TK|LM|DEN)-(\d+)\b', str(value)):
                maximum[patterns[prefix]] = max(maximum[patterns[prefix]], int(digits))
        for key in maximum:
            value = counters.get(key, 0)
            if isinstance(value, bool) or not isinstance(value, int) or not 0 <= value < 10**12:
                raise StorageError('Contadores heredados inválidos.')
            maximum[key] = value
        ids, specifics = set(), set()
        for sale in history:
            if not isinstance(sale, dict):
                raise StorageError('Una venta heredada no es válida.')
            ident, specific = sale.get('id_ticket_global'), sale.get('id_ticket_especifico')
            if not isinstance(ident, str) or not re.fullmatch(r'TK-\d+', ident) or not isinstance(specific, str) or not re.fullmatch(r'(LM|DEN)-\d+', specific):
                raise StorageError('Identificadores heredados inválidos.')
            if ident in ids or specific in specifics:
                raise StorageError('Existen IDs duplicados en el historial. Concilie los datos antes de importar.')
            if not isinstance(sale.get('paciente'), dict) or not isinstance(sale.get('items'), list) or sale.get('destino') not in ('General', 'Dental'):
                raise StorageError('Una venta heredada tiene campos incompletos.')
            money(sale.get('total_final'), 'Total histórico')
            ids.add(ident); specifics.add(specific)
            reserve(ident); reserve(specific)
        # Excel may contain older sales absent from JSON. Preserve its rows and reserve IDs.
        from openpyxl import load_workbook
        for name in ('Contabilidad_Las_Marianas.xlsx', 'Contabilidad_Dental.xlsx'):
            path = self.data_dir / name
            if path.exists():
                wb = load_workbook(path, read_only=True)
                try:
                    for row in wb.active.iter_rows(min_row=2, min_col=2, max_col=2, values_only=True):
                        reserve(row[0])
                finally:
                    wb.close()
        if self.tickets_dir.exists():
            for path in self.tickets_dir.glob('*.pdf'):
                reserve(path.stem.replace('_', ' '))
        legacy_files = [p for p in self.data_dir.iterdir() if p.is_file() and p.suffix in ('.json', '.xlsx')]
        if legacy_files:
            backup = self.data_dir / 'backups' / ('pre-migration-' + now().strftime('%Y%m%d-%H%M%S-') + uuid.uuid4().hex[:8])
            backup.mkdir(parents=True)
            for path in legacy_files:
                shutil.copy2(path, backup / path.name)
            db.execute("INSERT INTO metadata VALUES ('legacy_backup',?)", (str(backup),))
        for sale in history:
            sale.setdefault('habilitado', True)
            db.execute('INSERT INTO tickets VALUES (?,?,?,?,?,?)', (
                sale['id_ticket_global'], sale['id_ticket_especifico'], 'legacy:' + sale['id_ticket_global'],
                request_hash(sale), encode(sale), sale.get('fecha_iso', sale.get('fecha', ''))))
            self._event(db, sale['id_ticket_global'], 'imported', {'source': 'ventas.json', 'costs': 'historical values only'})
        for key, value in maximum.items():
            db.execute('INSERT INTO counters VALUES (?,?)', (key, value))

    @staticmethod
    def _event(db, ticket, kind, details=None):
        db.execute('INSERT INTO events(ticket_id,kind,at,details) VALUES (?,?,?,?)',
                   (ticket, kind, now().isoformat(), encode(details or {})))

    def existing(self, key, fingerprint):
        with self.connection() as db:
            row = db.execute('SELECT payload,request_hash FROM tickets WHERE operation_key=?', (key,)).fetchone()
        if row:
            if row['request_hash'] != fingerprint:
                raise ConflictError('La operación ya existe con otros datos. Consulte el historial.')
            return json.loads(row['payload'])
        return None

    def register(self, sale, key, fingerprint, copies=2):
        with self.connection() as db:
            db.execute('BEGIN IMMEDIATE')
            row = db.execute('SELECT payload,request_hash FROM tickets WHERE operation_key=?', (key,)).fetchone()
            if row:
                if row['request_hash'] != fingerprint:
                    raise ConflictError('La operación ya existe con otros datos.')
                return json.loads(row['payload']), False
            destination = sale['destino'].lower()
            for name in ('global', destination):
                db.execute('UPDATE counters SET value=value+1 WHERE name=?', (name,))
            values = dict(db.execute('SELECT name,value FROM counters').fetchall())
            sale = dict(sale, id_ticket_global=f"TK-{values['global']:04d}",
                        id_ticket_especifico=f"{'LM' if destination == 'general' else 'DEN'}-{values[destination]:04d}")
            ident = sale['id_ticket_global']
            db.execute('INSERT INTO tickets VALUES (?,?,?,?,?,?)', (ident, sale['id_ticket_especifico'], key, fingerprint, encode(sale), sale['fecha_iso']))
            self._event(db, ident, 'registered')
            for kind in ['excel', 'pdf'] + ['print'] * copies:
                db.execute('INSERT INTO jobs(ticket_id,kind) VALUES (?,?)', (ident, kind))
            return sale, True

    def get(self, ident):
        with self.connection() as db:
            row = db.execute('SELECT payload FROM tickets WHERE id=?', (ident,)).fetchone()
        return json.loads(row['payload']) if row else None

    def history(self):
        with self.connection() as db:
            rows = db.execute('SELECT payload FROM tickets ORDER BY rowid').fetchall()
        return [json.loads(row['payload']) for row in rows]

    def cancel(self, ident, reason):
        with self.connection() as db:
            db.execute('BEGIN IMMEDIATE')
            row = db.execute('SELECT payload FROM tickets WHERE id=?', (ident,)).fetchone()
            if row is None:
                return False
            sale = json.loads(row['payload'])
            if sale.get('habilitado') is False:
                return True
            sale.update(habilitado=False, motivo_eliminacion=reason, fecha_anulacion=now().isoformat())
            db.execute('UPDATE tickets SET payload=? WHERE id=?', (encode(sale), ident))
            db.execute("UPDATE jobs SET state='cancelled' WHERE ticket_id=? AND kind='print' AND state='pending'", (ident,))
            db.execute("INSERT INTO jobs(ticket_id,kind) VALUES (?,'excel')", (ident,))
            self._event(db, ident, 'cancelled', {'reason': reason})
            return True

    def jobs(self, ident):
        with self.connection() as db:
            return [dict(r) for r in db.execute('SELECT id,kind,state,error FROM jobs WHERE ticket_id=? ORDER BY id', (ident,))]

    def queue_print(self, ident, copies):
        with self.connection() as db:
            db.execute('BEGIN IMMEDIATE')
            if db.execute("SELECT 1 FROM jobs WHERE ticket_id=? AND kind='print' AND state IN ('pending','running')", (ident,)).fetchone():
                raise ConflictError('Hay una impresión pendiente para este ticket.')
            for _ in range(copies):
                db.execute("INSERT INTO jobs(ticket_id,kind) VALUES (?,'print')", (ident,))
            self._event(db, ident, 'reprint_requested', {'copies': copies})

    def retry_documents(self, ident):
        with self.connection() as db:
            db.execute("UPDATE jobs SET state='pending',error=NULL WHERE ticket_id=? AND state='failed' AND kind!='print'", (ident,))

    def claim_job(self):
        with self.connection() as db:
            db.execute('BEGIN IMMEDIATE')
            row = db.execute("SELECT * FROM jobs WHERE state='pending' ORDER BY id LIMIT 1").fetchone()
            if row:
                db.execute("UPDATE jobs SET state='running',error=NULL WHERE id=?", (row['id'],))
                return dict(row)
        return None

    def finish_job(self, job, success, error=None):
        with self.connection() as db:
            db.execute('UPDATE jobs SET state=?,error=? WHERE id=?', ('sent' if success else 'failed', error, job['id']))
            self._event(db, job['ticket_id'], job['kind'] + ('_sent' if success else '_failed'))

    def backup(self, target):
        """SQLite backup API creates a consistent snapshot, including uncheckpointed pages."""
        target = Path(target)
        if target.exists():
            raise StorageError('El destino de respaldo ya existe.')
        target.mkdir(parents=True)
        with self.connection() as db:
            with closing(sqlite3.connect(target / self.path.name)) as copy:
                db.backup(copy)
                if copy.execute('PRAGMA integrity_check').fetchone()[0] != 'ok':
                    raise StorageError('El respaldo no pasó la comprobación de integridad.')
        for path in self.data_dir.glob('*.xlsx'):
            shutil.copy2(path, target / path.name)
        if self.tickets_dir.exists():
            shutil.copytree(self.tickets_dir, target / 'Tickets')
        (target / 'BACKUP_COMPLETE').write_text(now().isoformat(), encoding='utf-8')
        return target
