"""Validated Excel catalogs and idempotent accounting exports (not primary storage)."""
from decimal import Decimal
import os
from pathlib import Path
import tempfile

from openpyxl import Workbook, load_workbook
from models import ValidationError, money, text

GENERAL = 'Lista de Precios-Las Marianas.xlsx'
DENTAL = 'Lista de Precios-Dental.xlsx'
COLUMNS = ['Fecha', 'nota de venta', 'Pago', 'Nombre', 'Descripción', 'A cuenta',
           'ingresos', 'Egresos - gastos', 'LABORATORIO', 'Insumos de lab',
           'Ganancia', 'Saldos', 'observaciones']

class CatalogError(RuntimeError):
    pass

def normalize_accounting_header(value):
    """Accept the existing note-number label without changing workbook headers."""
    normalized = ' '.join(str(value or '').split()).casefold()
    if normalized == 'id nota de venta':
        return 'nota de venta'
    return normalized


def cargar_catalogos(data_dir):
    result = {}
    for destination, filename, ident in [('general', GENERAL, 'ID_General'), ('dental', DENTAL, 'ID_Dental')]:
        path = Path(data_dir) / filename
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
            try:
                rows = iter(wb.active.values)
                headers = next(rows, ())
                required = {ident, 'Categoría', 'Nombre Específico', 'Precio Unitario'}
                required |= {'Costo Laboratorio', 'Costo Insumos'} if destination == 'general' else {'Costo'}
                if len(headers) != len(set(headers)) or not required.issubset(headers):
                    raise CatalogError(f'{filename}: faltan columnas o existen encabezados duplicados.')
                records, ids = [], set()
                for values in rows:
                    if not any(v is not None for v in values):
                        continue
                    row = dict(zip(headers, values))
                    service_id = text(str(row.get(ident) or ''), 'ID de catálogo', 80)
                    if service_id in ids or service_id == 'OTROS':
                        raise CatalogError(f'{filename}: ID duplicado o reservado.')
                    ids.add(service_id)
                    record = {ident: service_id, 'Categoría': text(row['Categoría'], 'Categoría'),
                              'Nombre Específico': text(row['Nombre Específico'], 'Servicio')}
                    for column in ['Precio Unitario', 'Descuento'] + (['Costo Laboratorio', 'Costo Insumos'] if destination == 'general' else ['Costo']):
                        record[column] = str(money(row.get(column) if row.get(column) is not None else 0, column))
                    if Decimal(record['Descuento']) > Decimal(record['Precio Unitario']):
                        raise CatalogError(f'{filename}: un descuento supera el precio.')
                    records.append(record)
                result[destination] = records
            finally:
                wb.close()
        except (OSError, ValueError, KeyError, TypeError) as exc:
            raise CatalogError(f'No se pudo cargar {filename}. Revise el archivo y sus columnas.') from exc
    return result

def atomic_bytes(path, contents):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(prefix='.ticket-', suffix=path.suffix, dir=path.parent)
    try:
        with os.fdopen(fd, 'wb') as stream:
            stream.write(contents)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)

def sync_accounting(data_dir, sale):
    filename = 'Contabilidad_Dental.xlsx' if sale['destino'] == 'Dental' else 'Contabilidad_Las_Marianas.xlsx'
    path = Path(data_dir) / filename
    wb = load_workbook(path) if path.exists() else Workbook()
    try:
        ws = wb.active
        if not path.exists():
            ws.append(COLUMNS)
        headers = list(next(ws.values))[:len(COLUMNS)]
        if [normalize_accounting_header(value) for value in headers] != [normalize_accounting_header(value) for value in COLUMNS]:
            raise CatalogError('El libro contable tiene encabezados incompatibles. Revise sus columnas; el archivo se conservó sin cambios.')
        matches = [r for r in range(2, ws.max_row + 1)
                   if str(ws.cell(r, 2).value or '').split('|', 1)[0].strip() == sale['id_ticket_global']]
        if sale.get('habilitado') is False:
            for row in reversed(matches):
                ws.delete_rows(row)
        else:
            # Unknown historical costs must never be invented on re-export.
            if any('costo_unitario' not in item for item in sale['items']):
                raise CatalogError('La venta histórica no conserva costos; requiere conciliación manual.')
            income = Decimal(sale['total_final'])
            lab, supplies = Decimal('0'), Decimal('0')
            for item in sale['items']:
                count = item['cantidad']
                if sale['destino'] == 'General':
                    lab += Decimal(item['costo_lab']) * count
                    supplies += Decimal(item['costo_insumo']) * count
                elif any(word in item['categoria'].lower() for word in ('insumo', 'medicamento')):
                    supplies += Decimal(item['costo_unitario']) * count
                else:
                    lab += Decimal(item['costo_unitario']) * count
            observation = sale.get('observaciones', '')
            if sale.get('descuento_especial_activo'):
                observation += ' | Desc: ' + sale['descuento_especial_razon']
            values = [sale['fecha'].split(' ')[-1].replace('/', '.'),
                      f"{sale['id_ticket_global']} | {sale['id_ticket_especifico']}", sale['metodo_pago'],
                      sale['paciente']['nombre'], ', '.join(f"{i['nombre']} x{i['cantidad']}" for i in sale['items']),
                      0, float(income), 0, float(lab), float(supplies), float(income - lab - supplies), '', observation]
            row = matches[0] if matches else ws.max_row + 1
            for duplicate in reversed(matches[1:]):
                ws.delete_rows(duplicate)
            for column, value in enumerate(values, 1):
                cell = ws.cell(row, column, value)
                if isinstance(value, str):
                    cell.data_type = 's'  # user text beginning '=' is never an Excel formula
        for row in range(2, ws.max_row + 1):
            ws.cell(row, 12, '=K2' if row == 2 else f'=L{row-1}+K{row}')
        from io import BytesIO
        output = BytesIO()
        wb.save(output)
        atomic_bytes(path, output.getvalue())
    finally:
        wb.close()
