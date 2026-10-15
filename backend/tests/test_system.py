import copy
from concurrent.futures import ThreadPoolExecutor
import json
from pathlib import Path
import sqlite3
import subprocess
import sys
import tempfile
import threading
import urllib.request
import unittest
from unittest.mock import patch
import uuid

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'backend'))
sys.path.insert(0, str(ROOT))
from app import create_app
from catalog_manager import sync_accounting, cargar_catalogos, CatalogError
from generar_precios import generate
from models import build_sale, money, ValidationError
from pdf_generator import generar_ticket_pdf
from receipt import receipt_lines
from storage import Store, StorageError
from openpyxl import Workbook, load_workbook


class SystemTests(unittest.TestCase):
    def setUp(self):
        temporary_root = ROOT / '.test-runtime'
        temporary_root.mkdir(exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(dir=temporary_root)
        self.base = Path(self.temp.name)
        self.data = generate(self.base / 'data')
        self.tickets = self.base / 'Tickets'
        self.app = create_app({'TESTING': True, 'DATA_DIR': str(self.data), 'TICKETS_DIR': str(self.tickets),
                               'START_WORKER': False, 'AUTO_BACKUP': False, 'PRINTER': lambda _: False})
        self.client = self.app.test_client()
        self.body = {'paciente': {'nombre': 'PACIENTE FICTICIO', 'dni': '12345678'}, 'destino': 'general',
                     'items': [{'id': 'LM-001', 'nombre': 'TEXTO NO CONFIABLE', 'cantidad': 1}],
                     'desglose_pagos': [{'metodo': 'Efectivo', 'monto': '30.00'}], 'total_esperado': '30.00'}

    def tearDown(self):
        self.app.extensions['worker'].stop()
        # TemporaryDirectory is rooted explicitly inside the authorized workspace.
        self.temp.cleanup()

    @property
    def store(self):
        return self.app.extensions['store']

    def post(self, body=None, key=None):
        return self.client.post('/api/registrar-venta', json=body if body is not None else self.body,
                                headers={'X-Local-Request': '1', 'Idempotency-Key': key or str(uuid.uuid4())})

    def drain(self):
        while self.app.extensions['worker'].process_one():
            pass

    def test_registration_replay_survives_catalog_change(self):
        key = str(uuid.uuid4())
        first = self.post(key=key)
        self.assertEqual(first.status_code, 200)
        path = self.data / 'Lista de Precios-Las Marianas.xlsx'
        wb = load_workbook(path); wb.active.cell(2, 4, 40); wb.save(path); wb.close()
        replay = self.post(key=key)
        self.assertEqual(replay.json['ticket_id'], first.json['ticket_id'])
        self.assertTrue(replay.json['replayed'])
        self.assertEqual(len(self.store.history()), 1)
        self.assertEqual(len(self.store.jobs(first.json['ticket_id'])), 4)
        self.assertEqual(self.store.history()[0]['items'][0]['nombre'], 'Consulta Médica General')

    def test_operation_key_cannot_be_reused_with_changed_data(self):
        key = str(uuid.uuid4())
        self.post(key=key)
        changed = dict(self.body, observaciones='Otra operación')
        self.assertEqual(self.post(changed, key).status_code, 409)
        self.assertEqual(len(self.store.history()), 1)

    def test_invalid_sales_do_not_consume_ids(self):
        cases = []
        for count in [-1, 0, 1.5, True, 1000]:
            body = copy.deepcopy(self.body); body['items'][0]['cantidad'] = count; cases.append(body)
        for dni in ['123', 'abcdefgh']:
            body = copy.deepcopy(self.body); body['paciente']['dni'] = dni; cases.append(body)
        cases += [dict(self.body, destino='otro'), dict(self.body, items=[]), dict(self.body, desglose_pagos=[]),
                  dict(self.body, descuento_especial_activo='false'),
                  dict(self.body, descuento_especial_activo=True, descuento_especial_tipo='percent', descuento_especial_valor=101, descuento_especial_razon='Prueba'),
                  dict(self.body, desglose_pagos=[{'metodo': 'Yape', 'monto': '40.00'}]),
                  dict(self.body, total_esperado='29.99')]
        for body in cases:
            with self.subTest(body=body):
                self.assertEqual(self.post(body).status_code, 422)
        result = self.post()
        self.assertEqual(result.json['ticket_id'], 'TK-0001')

    def test_inactive_discount_and_client_change_are_not_trusted(self):
        body = dict(self.body, descuento_especial_activo=False, descuento_especial_monto_soles='30', vuelto='9999', pago_con='0')
        result = self.post(body)
        self.assertEqual(result.status_code, 200)
        sale = self.store.history()[0]
        self.assertEqual(sale['total_final'], '30.00')
        self.assertEqual(sale['vuelto'], '0.00')
        self.assertEqual(sale['pago_con'], '30.00')

    def test_catalog_discount_and_half_up_rounding(self):
        catalog = cargar_catalogos(self.data)
        catalog['general'][0]['Descuento'] = '0.01'
        body = dict(self.body, total_esperado='27.99', descuento_especial_activo=True,
                    descuento_especial_tipo='fixed', descuento_especial_valor='2', descuento_especial_razon='Convenio')
        sale = build_sale(body, catalog)
        self.assertEqual(sale['total_final'], '27.99')
        self.assertEqual(sale['vuelto'], '2.01')

    def test_currency_rejects_ambiguous_and_nonfinite_values(self):
        for value in ['-1', 'NaN', 'Infinity', '1,000.00', 'S/30', '1.001', True]:
            with self.subTest(value=value), self.assertRaises(ValidationError):
                money(value)

    def test_concurrent_retries_create_one_sale(self):
        self.app.extensions['prepare']()
        key = str(uuid.uuid4())
        def submit(_):
            with self.app.test_client() as client:
                return client.post('/api/registrar-venta', json=self.body,
                    headers={'X-Local-Request': '1', 'Idempotency-Key': key}).status_code
        with ThreadPoolExecutor(max_workers=6) as pool:
            results = list(pool.map(submit, range(12)))
        self.assertEqual(results, [200] * 12)
        self.assertEqual(len(self.store.history()), 1)
        self.assertEqual(len(self.store.jobs('TK-0001')), 4)

    def test_concurrent_distinct_sales_receive_unique_ids(self):
        self.app.extensions['prepare']()
        def submit(_):
            with self.app.test_client() as client:
                response = client.post('/api/registrar-venta', json=self.body,
                    headers={'X-Local-Request': '1', 'Idempotency-Key': str(uuid.uuid4())})
                self.assertEqual(response.status_code, 200)
                return response.json['ticket_id']
        with ThreadPoolExecutor(max_workers=6) as pool:
            ids = list(pool.map(submit, range(12)))
        self.assertEqual(len(set(ids)), 12)

    def test_database_failure_rolls_back_sale_counter_and_jobs(self):
        self.app.extensions['prepare']()
        with self.store.connection() as db:
            db.execute("CREATE TRIGGER fail_job BEFORE INSERT ON jobs BEGIN SELECT RAISE(ABORT,'synthetic'); END")
        self.assertEqual(self.post().status_code, 500)
        self.assertEqual(self.store.history(), [])
        with self.store.connection() as db:
            self.assertEqual(db.execute("SELECT value FROM counters WHERE name='global'").fetchone()[0], 0)
            db.execute('DROP TRIGGER fail_job')
        self.assertEqual(self.post().json['ticket_id'], 'TK-0001')

    def test_pdf_and_print_failures_do_not_undo_sale(self):
        self.assertEqual(self.post().status_code, 200)
        with patch('jobs.generar_ticket_pdf', side_effect=ValueError('synthetic PDF error')):
            self.drain()
        states = self.store.jobs('TK-0001')
        self.assertEqual([(j['kind'], j['state']) for j in states], [('excel', 'sent'), ('pdf', 'failed'), ('print', 'failed'), ('print', 'failed')])
        self.assertEqual(len(self.store.history()), 1)
        self.client.post('/api/tickets/TK-0001/reintentar-documentos', headers={'X-Local-Request': '1'})
        self.drain()
        self.assertTrue((self.tickets / 'TK-0001.pdf').exists())
        self.assertEqual(sum(j['state'] == 'failed' for j in self.store.jobs('TK-0001')), 2)

    def test_reprint_limits_and_pending_protection(self):
        self.post()
        for count in [0, -1, 6, 1.5, True]:
            result = self.client.post('/api/reimprimir-ticket', json={'id_ticket_global': 'TK-0001', 'cantidad': count}, headers={'X-Local-Request': '1'})
            self.assertEqual(result.status_code, 422)
        result = self.client.post('/api/reimprimir-ticket', json={'id_ticket_global': 'TK-0001', 'cantidad': 1}, headers={'X-Local-Request': '1'})
        self.assertEqual(result.status_code, 409)

    def test_interrupted_print_is_not_automatically_repeated(self):
        self.post()
        self.app.extensions['worker'].process_one()
        self.app.extensions['worker'].process_one()
        claimed = self.store.claim_job()
        self.assertEqual(claimed['kind'], 'print')
        restarted = Store(self.data, self.tickets); restarted.initialize()
        self.assertEqual(restarted.jobs('TK-0001')[2]['state'], 'uncertain')

    def test_cancel_is_idempotent_and_requires_reason(self):
        self.post(); self.drain()
        headers = {'X-Local-Request': '1'}
        self.assertEqual(self.client.delete('/api/eliminar-ticket/TK-0001', json={'motivo': ''}, headers=headers).status_code, 422)
        for _ in range(2):
            self.assertEqual(self.client.delete('/api/eliminar-ticket/TK-0001', json={'motivo': 'Cancelación ficticia'}, headers=headers).status_code, 200)
        self.drain()
        self.assertFalse(self.store.get('TK-0001')['habilitado'])
        with self.store.connection() as db:
            self.assertEqual(db.execute("SELECT COUNT(*) FROM events WHERE kind='cancelled'").fetchone()[0], 1)
        self.assertTrue((self.tickets / 'TK-0001.pdf').exists())
        wb = load_workbook(self.data / 'Contabilidad_Las_Marianas.xlsx'); self.assertEqual(wb.active.max_row, 1); wb.close()

    def test_accounting_exact_id_safe_text_and_idempotent_export(self):
        sale = build_sale(self.body, cargar_catalogos(self.data))
        sale.update(id_ticket_global='TK-10000', id_ticket_especifico='LM-10000', observaciones='=1+1')
        sync_accounting(self.data, sale); sync_accounting(self.data, sale)
        different = dict(sale, id_ticket_global='TK-1000', id_ticket_especifico='LM-1000', habilitado=False)
        sync_accounting(self.data, different)
        wb = load_workbook(self.data / 'Contabilidad_Las_Marianas.xlsx'); ws = wb.active
        self.assertEqual(ws.max_row, 2)
        self.assertEqual(ws.cell(2, 13).data_type, 's')
        self.assertEqual(ws.cell(2, 13).value, '=1+1'); wb.close()

    def test_locked_export_leaves_original_intact(self):
        sale = build_sale(self.body, cargar_catalogos(self.data))
        sale.update(id_ticket_global='TK-0001', id_ticket_especifico='LM-0001')
        sync_accounting(self.data, sale)
        path = self.data / 'Contabilidad_Las_Marianas.xlsx'; original = path.read_bytes()
        with patch('catalog_manager.os.replace', side_effect=PermissionError('locked')):
            with self.assertRaises(PermissionError):
                sync_accounting(self.data, dict(sale, habilitado=False))
        self.assertEqual(path.read_bytes(), original)

    def test_corrupt_legacy_history_is_not_overwritten(self):
        path = self.data / 'ventas.json'; original = b'{broken JSON'; path.write_bytes(original)
        response = self.post()
        self.assertEqual(response.status_code, 503)
        self.assertEqual(path.read_bytes(), original)
        with self.store.connection() as db:
            self.assertEqual(db.execute('SELECT COUNT(*) FROM tickets').fetchone()[0], 0)

    def test_legacy_import_reserves_excel_ids_and_backs_up(self):
        old = build_sale(self.body, cargar_catalogos(self.data))
        old.update(id_ticket_global='TK-0001', id_ticket_especifico='LM-0001', habilitado=False)
        path = self.data / 'ventas.json'; path.write_text(json.dumps([old]), encoding='utf-8')
        original = path.read_bytes()
        wb = Workbook(); wb.active.append(['Fecha', 'nota de venta']); wb.active.append(['test', 'TK-10000 | LM-10000']); wb.save(self.data / 'Contabilidad_Las_Marianas.xlsx'); wb.close()
        response = self.post()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json['ticket_id'], 'TK-10001')
        self.assertEqual(path.read_bytes(), original)
        self.assertFalse(self.store.get('TK-0001')['habilitado'])
        self.assertEqual(len(list((self.data / 'backups').glob('pre-migration-*/ventas.json'))), 1)

    def test_backup_restores_history_and_job_state(self):
        self.post()
        target = self.store.backup(self.base / 'backup')
        restored = Store(target, target / 'Tickets'); restored.initialize()
        self.assertEqual(restored.history(), self.store.history())
        self.assertEqual(len(restored.jobs('TK-0001')), 4)
        self.assertTrue((target / 'BACKUP_COMPLETE').exists())

    def test_restore_command_marks_old_prints_uncertain(self):
        self.post()
        backup = self.store.backup(self.base / 'backup')
        destination = self.base / 'restored'
        result = subprocess.run([sys.executable, '-B', str(ROOT / 'backend/maintenance.py'),
                                 'restore', '--source', str(backup), '--output', str(destination)],
                                capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        restored = Store(destination, destination / 'Tickets'); restored.initialize()
        self.assertEqual([j['state'] for j in restored.jobs('TK-0001')], ['pending', 'pending', 'uncertain', 'uncertain'])
        self.assertEqual([j['state'] for j in self.store.jobs('TK-0001')], ['pending'] * 4)

    def test_backup_command_does_not_reset_live_job(self):
        self.post()
        claimed = self.store.claim_job()
        result = subprocess.run([sys.executable, '-B', str(ROOT / 'backend/maintenance.py'),
                                 '--data-dir', str(self.data), '--tickets-dir', str(self.tickets),
                                 'backup', '--output', str(self.base / 'backup')], capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(self.store.jobs('TK-0001')[0]['state'], 'running')
        self.store.finish_job(claimed, True)

    def test_waitress_serves_local_api(self):
        from waitress import create_server
        server = create_server(self.app, host='127.0.0.1', port=0, map={})
        thread = threading.Thread(target=server.run, daemon=True)
        thread.start()
        try:
            url = f'http://127.0.0.1:{server.effective_port}'
            with urllib.request.urlopen(url + '/api/health', timeout=5) as response:
                self.assertEqual(json.load(response)['status'], 'ok')
            request = urllib.request.Request(url + '/api/registrar-venta', data=json.dumps(self.body).encode(),
                headers={'Content-Type': 'application/json', 'X-Local-Request': '1', 'Idempotency-Key': str(uuid.uuid4()), 'Origin': url})
            with urllib.request.urlopen(request, timeout=5) as response:
                self.assertEqual(json.load(response)['ticket_id'], 'TK-0001')
        finally:
            server.close()
            server.task_dispatcher.shutdown()
            thread.join(timeout=5)

    def test_catalog_duplicate_ids_and_demo_overwrite_are_rejected(self):
        with self.assertRaises(FileExistsError): generate(self.data)
        path = self.data / 'Lista de Precios-Las Marianas.xlsx'
        wb = load_workbook(path); wb.active.cell(3, 1, 'LM-001'); wb.save(path); wb.close()
        with self.assertRaises(CatalogError): cargar_catalogos(self.data)

    def test_receipt_preserves_long_text_unicode_and_observations(self):
        sale = build_sale(dict(self.body, observaciones='Observación con Ñ y texto de prueba.'), cargar_catalogos(self.data))
        sale.update(id_ticket_global='TK-0001', id_ticket_especifico='LM-0001')
        sale['items'][0]['nombre'] = 'DESCRIPCIÓN COMPLETA ' * 8
        content = ' '.join(receipt_lines(sale))
        self.assertIn('Observación', content)
        self.assertEqual(content.count('DESCRIPCIÓN'), 8)
        output = generar_ticket_pdf(sale)
        self.assertTrue(output.startswith(b'%PDF-'))
        self.assertIn(b'/MediaBox', output)

    def test_local_origin_header_and_unknown_api_protection(self):
        self.assertEqual(self.client.post('/api/registrar-venta', json=self.body).status_code, 403)
        self.assertEqual(self.client.get('/api/tickets', headers={'Host': 'attacker.example'}).status_code, 403)
        self.assertEqual(self.client.get('/api/tickets', headers={'Origin': 'https://attacker.example'}).status_code, 403)
        self.assertEqual(self.client.get('/api/does-not-exist').status_code, 404)
        self.assertEqual(self.client.get('/api/health').json['storage'], 'sqlite')


if __name__ == '__main__':
    unittest.main()
